import io
import logging
from datetime import datetime, date, timezone
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func

from app.api.dependencies.deps import get_db, get_current_user
from app.models.user import User
from app.models.workday import Workday
from app.models.visit import Visit
from app.models.report import Report
from app.models.complaint import Complaint

logger = logging.getLogger(__name__)

try:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    HAS_REPORTLAB = True
except ImportError:
    HAS_REPORTLAB = False
    logger.warning("Reportlab is not installed. PDF exports will fail.")

from sqlalchemy.orm import Session, joinedload

router = APIRouter()

@router.get("/workdays/all")
def get_all_workdays_stats(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    if current_user.role not in ["admin", "supervisor"]:
        raise HTTPException(status_code=403, detail="Forbidden")
    
    # Get all workdays, joined with user, visits and reports
    workdays = db.query(Workday).options(
        joinedload(Workday.user),
        joinedload(Workday.visits),
        joinedload(Workday.reports)
    ).order_by(Workday.start_time.desc()).all()
    
    results = []
    for wd in workdays:
        # Calculate duration
        duration_str = "Active"
        if wd.start_time and wd.end_time:
            diff = wd.end_time - wd.start_time
            hrs = int(diff.total_seconds() // 3600)
            mins = int((diff.total_seconds() % 3600) // 60)
            duration_str = f"{hrs}h {mins}m"
        
        results.append({
            "id": wd.id,
            "merchandiser": f"{wd.user.first_name} {wd.user.last_name}" if wd.user else "Unknown",
            "date": wd.start_time.isoformat() if wd.start_time else None,
            "stores_visited": len(wd.visits),
            "reports_submitted": len(wd.reports),
            "duration": duration_str,
            "status": wd.status,
            "user_id": wd.user_id,
            "user_email": wd.user.email if wd.user else "---",
            "user_role": wd.user.role if wd.user else "---"
        })
    
    return results

import os
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image

@router.get("/workday/{workday_id}/pdf")
def export_workday_pdf(
    workday_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    if not HAS_REPORTLAB:
        raise HTTPException(status_code=500, detail="PDF generation library not installed.")
        
    if current_user.role not in ["admin", "supervisor"]:
        raise HTTPException(status_code=403, detail="Forbidden")

    workday = db.query(Workday).options(
        joinedload(Workday.user),
        joinedload(Workday.visits).joinedload(Visit.gms)
    ).filter(Workday.id == workday_id).first()
    
    if not workday:
        raise HTTPException(status_code=404, detail="Workday not found")

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
    story = []
    styles = getSampleStyleSheet()
    
    # Header
    title_style = styles['Heading1']
    title_style.alignment = 1 # Center
    story.append(Paragraph(f"Detailed Workday Report", title_style))
    story.append(Spacer(1, 20))
    
    # Metadata Table
    meta_data = [
        ["Merchandiser", f"{workday.user.first_name} {workday.user.last_name}"],
        ["Date", workday.start_time.strftime('%B %d, %Y')],
        ["Shift", f"{workday.start_time.strftime('%H:%M')} - {workday.end_time.strftime('%H:%M') if workday.end_time else 'Active'}"],
        ["Status", workday.status.upper()]
    ]
    mt = Table(meta_data, colWidths=[120, 300])
    mt.setStyle(TableStyle([
        ('FONTNAME', (0,0), (0,-1), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 8),
        ('GRID', (0,0), (-1,-1), 0.5, colors.grey),
    ]))
    story.append(mt)
    story.append(Spacer(1, 24))

    # Visits Table
    story.append(Paragraph("Store Visits Summary", styles['Heading2']))
    story.append(Spacer(1, 10))
    visit_data = [["Store", "Arrival", "Departure", "Duration"]]
    
    for v in workday.visits:
        duration_str = "In Progress"
        if v.end_time:
            diff = v.end_time - v.start_time
            mins = int(diff.total_seconds() / 60)
            duration_str = f"{mins} min"
            
        visit_data.append([
            v.gms.name if v.gms else f"Store ID: {v.gms_id}",
            v.start_time.strftime("%H:%M"),
            v.end_time.strftime("%H:%M") if v.end_time else "-",
            duration_str
        ])

    vt = Table(visit_data, colWidths=[180, 80, 80, 80])
    vt.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.whitesmoke, colors.beige])
    ]))
    story.append(vt)
    story.append(Spacer(1, 30))

    # Group reports by visit for organized display
    # We gather reports directly linked to workday AND those linked via visits
    visit_ids = [v.id for v in workday.visits]
    all_reports = db.query(Report).filter(
        (Report.workday_id == workday_id) | (Report.visit_id.in_(visit_ids))
    ).all()

    reports_by_visit = {}
    standalone_reports = []
    
    for r in all_reports:
        if r.visit_id and r.visit_id in visit_ids:
            if r.visit_id not in reports_by_visit: reports_by_visit[r.visit_id] = []
            reports_by_visit[r.visit_id].append(r)
        else:
            standalone_reports.append(r)

    # Detailed Task & Photos Section
    story.append(Paragraph("Timeline & Field Activity Details", styles['Heading2']))
    story.append(Spacer(1, 12))
    
    if not workday.visits:
        story.append(Paragraph("No store visits recorded for this shift.", styles['Normal']))
    
    for v in workday.visits:
        # Visit Header Card-like display
        store_name = v.gms.name if v.gms else f"Store #{v.gms_id}"
        story.append(Paragraph(f"<font color='blue' size=14><b>{store_name}</b></font>", styles['Normal']))
        story.append(Paragraph(f"Visit Time: {v.start_time.strftime('%H:%M')} - {v.end_time.strftime('%H:%M') if v.end_time else 'Active'}", styles['Normal']))
        story.append(Spacer(1, 8))

        visit_reports = reports_by_visit.get(v.id, [])
        if not visit_reports:
            story.append(Paragraph("<i>No specific reports submitted for this visit.</i>", styles['Normal']))
        else:
            for r in visit_reports:
                # Highlight Anomalies
                is_anomaly = r.type.lower() == "anomaly"
                header_color = "red" if is_anomaly else "black"
                label = "[ANOMALY] " if is_anomaly else ""
                
                story.append(Paragraph(f"<font color='{header_color}'><b>{label}{r.name}</b></font>", styles['Normal']))
                if r.notes:
                    story.append(Paragraph(f"Description: {r.notes}", styles['Normal']))
                
                # Photos
                row_labels = []
                row_imgs = []
                
                def get_img_path(url):
                    if not url or url.startswith("blob:"): return None
                    
                    idx = url.rfind("/static/")
                    if idx != -1:
                        rel_path = url[idx + 8:].lstrip("/") # 8 is len of "/static/"
                        full_path = os.path.join("uploads", rel_path)
                        return full_path if os.path.exists(full_path) else None
                    return None

                if r.before_image:
                    p = get_img_path(r.before_image)
                    if p:
                        try:
                            row_labels.append(Paragraph("BEFORE", styles['Normal']))
                            row_imgs.append(Image(p, width=180, height=180))
                        except: pass
                
                if r.after_image:
                    p = get_img_path(r.after_image)
                    if p:
                        try:
                            row_labels.append(Paragraph("AFTER", styles['Normal']))
                            row_imgs.append(Image(p, width=180, height=180))
                        except: pass

                if row_imgs:
                    num_imgs = len(row_imgs)
                    img_table = Table([row_labels, row_imgs], colWidths=[200] * num_imgs)
                    img_table.setStyle(TableStyle([
                        ('ALIGN', (0,0), (-1,-1), 'LEFT'),
                        ('TOPPADDING', (0,0), (-1,-1), 5),
                        ('BOTTOMPADDING', (0,0), (-1,-1), 10),
                    ]))
                    story.append(img_table)
                
                story.append(Spacer(1, 8))
        
        story.append(Spacer(1, 10))
        story.append(Paragraph("<font color='grey'>----------------------------------------------------------------------------------------------------------------------------------</font>", styles['Normal']))
        story.append(Spacer(1, 15))

    # Add standalone reports (shift summary etc)
    if standalone_reports:
        story.append(Paragraph("Shift Summaries & Other Logs", styles['Heading2']))
        for r in standalone_reports:
            story.append(Paragraph(f"<b>{r.name}</b>", styles['Normal']))
            story.append(Paragraph(f"{r.notes or ''}", styles['Normal']))
            story.append(Spacer(1, 10))

    doc.build(story)
    buffer.seek(0)
    
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename=detailed_report_{workday_id}.pdf"}
    )

@router.get("/daily-report")
def export_daily_report(
    target_date: date = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    if not HAS_REPORTLAB:
        raise HTTPException(
            status_code=500, 
            detail="PDF generation library (reportlab) is not installed on the server."
        )
        
    if current_user.role not in ["admin", "supervisor"]:
        raise HTTPException(status_code=403, detail="Forbidden: Admin or Supervisor access required.")

    # Use UTC for consistent date filtering across timezones
    if not target_date:
        target_date = datetime.now(timezone.utc).date()

    # Optimized Querying using count() for efficiency
    merch_count = db.query(User).filter(User.role == "merchandiser").count()
    active_workdays = db.query(Workday).filter(func.date(Workday.start_time) == target_date).count()
    
    visits_completed = db.query(Visit).join(Workday).filter(
        func.date(Workday.start_time) == target_date, 
        Visit.status == "completed"
    ).count()
    
    reports_submitted = db.query(Report).filter(func.date(Report.created_at) == target_date).count()
    complaints_count = db.query(Complaint).filter(func.date(Complaint.created_at) == target_date).count()

    # Create PDF in memory
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
    story = []
    
    styles = getSampleStyleSheet()
    title_style = styles['Heading1']
    title_style.alignment = 1 # Center
    
    story.append(Paragraph(f"Merchandising Daily Report", title_style))
    story.append(Paragraph(f"Generated on: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S')} UTC", styles['Normal']))
    story.append(Paragraph(f"Report Date: {target_date.strftime('%B %d, %Y')}", styles['Normal']))
    story.append(Spacer(1, 20))

    # Executive Summary Table
    story.append(Paragraph("Executive Summary", styles['Heading2']))
    story.append(Spacer(1, 10))
    
    summary_data = [
        ["Metric", "Value"],
        ["Total Merchandisers", str(merch_count)],
        ["Active Workdays Today", str(active_workdays)],
        ["Completed Visits", str(visits_completed)],
        ["Reports Submitted", str(reports_submitted)],
        ["New Complaints", str(complaints_count)],
    ]

    summary_table = Table(summary_data, colWidths=[200, 100])
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#4F46E5")),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.whitesmoke),
        ('GRID', (0, 0), (-1, -1), 1, colors.grey),
    ]))
    story.append(summary_table)
    story.append(Spacer(1, 30))

    # Detailed Activity Table
    story.append(Paragraph("Merchandiser Activity Detail", styles['Heading2']))
    story.append(Spacer(1, 10))

    workdays = db.query(Workday).filter(func.date(Workday.start_time) == target_date).all()
    if workdays:
        merch_detail_data = [["Merchandiser", "Start", "End", "Status"]]
        for wd in workdays:
            # Note: In production, consider joinedload(Workday.user) to avoid N+1
            user_name = f"{wd.user.first_name} {wd.user.last_name}" if wd.user else f"ID: {wd.user_id}"
            start = wd.start_time.strftime("%H:%M") if wd.start_time else "N/A"
            end = wd.end_time.strftime("%H:%M") if wd.end_time else "N/A"
            merch_detail_data.append([user_name, start, end, wd.status.capitalize()])
            
        detail_table = Table(merch_detail_data, colWidths=[150, 100, 100, 100])
        detail_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('GRID', (0, 0), (-1, -1), 1, colors.grey),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.whitesmoke, colors.beige])
        ]))
        story.append(detail_table)
    else:
        story.append(Paragraph("No merchandiser activity recorded for this date.", styles['Normal']))

    doc.build(story)
    
    buffer.seek(0)
    filename = f"daily_report_{target_date}.pdf"
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

@router.get("/daily-report/excel")
def export_daily_report_excel(
    target_date: date = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    if current_user.role not in ["admin", "supervisor"]:
        raise HTTPException(status_code=403, detail="Forbidden: Admin or Supervisor access required.")

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        raise HTTPException(status_code=500, detail="Excel generation library (openpyxl) is not installed.")

    if not target_date:
        target_date = datetime.now(timezone.utc).date()

    wb = Workbook()
    
    # --- Sheet 1: Executive Summary ---
    ws_summary = wb.active
    ws_summary.title = "Executive Summary"
    
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4F46E5")
    center_align = Alignment(horizontal="center", vertical="center")

    ws_summary.append(["Metric", "Value"])
    for cell in ws_summary[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    merch_count = db.query(User).filter(User.role == "merchandiser").count()
    active_workdays = db.query(Workday).filter(func.date(Workday.start_time) == target_date).count()
    visits_completed = db.query(Visit).join(Workday).filter(
        func.date(Workday.start_time) == target_date, 
        Visit.status == "completed"
    ).count()
    reports_submitted = db.query(Report).filter(func.date(Report.created_at) == target_date).count()
    complaints_count = db.query(Complaint).filter(func.date(Complaint.created_at) == target_date).count()

    ws_summary.append(["Total Merchandisers", merch_count])
    ws_summary.append(["Active Workdays Today", active_workdays])
    ws_summary.append(["Completed Visits", visits_completed])
    ws_summary.append(["Reports Submitted", reports_submitted])
    ws_summary.append(["New Complaints", complaints_count])
    
    ws_summary.column_dimensions['A'].width = 25
    ws_summary.column_dimensions['B'].width = 15

    # --- Sheet 2: Merchandiser Activity ---
    ws_activity = wb.create_sheet(title="Merchandiser Activity")
    ws_activity.append(["Merchandiser", "Start Time", "End Time", "Status"])
    
    for cell in ws_activity[1]:
        cell.font = header_font
        cell.fill = PatternFill("solid", fgColor="333333")
        cell.alignment = center_align

    workdays = db.query(Workday).filter(func.date(Workday.start_time) == target_date).all()
    for wd in workdays:
        user_name = f"{wd.user.first_name} {wd.user.last_name}" if wd.user else f"ID: {wd.user_id}"
        start = wd.start_time.strftime("%H:%M") if wd.start_time else "N/A"
        end = wd.end_time.strftime("%H:%M") if wd.end_time else "N/A"
        ws_activity.append([user_name, start, end, wd.status.capitalize()])

    ws_activity.column_dimensions['A'].width = 30
    ws_activity.column_dimensions['B'].width = 15
    ws_activity.column_dimensions['C'].width = 15
    ws_activity.column_dimensions['D'].width = 15

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"daily_report_{target_date}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )
